import base64
import io
import os
import unittest
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

try:
    import httpx2 as httpx
except ImportError:
    import httpx
import openai
from docx import Document
from backend.main import app
from backend.apply import TailoredResume


class ApplyTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test')

    async def asyncTearDown(self):
        await self.client.aclose()

    async def test_health_identifies_the_app_without_private_data(self):
        response = await self.client.get('/health')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'service':'jobmatch-ai', 'apiVersion':2, 'features':['apply-workspace','saved-profile','tailored-autofill','saved-autofill']})

    async def test_apply_and_extension_download(self):
        response = await self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Their next hire.', response.text)
        response = await self.client.get('/analysis')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Apply workspace', response.text)
        response = await self.client.get('/apply')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Their next hire.', response.text)
        response = await self.client.get('/extension.zip')
        with zipfile.ZipFile(io.BytesIO(response.content)) as z:
            self.assertEqual(set(z.namelist()), {'manifest.json','background.js','panel.html','panel.js','config.js','panel.css','job.js','fill.js','README.md'})

    async def test_document_round_trip(self):
        response = await self.client.post('/resume-document', json={'resume':'Jonathan Deng\nExperience\nBuilt APIs'})
        self.assertEqual(response.status_code, 200)
        doc = Document(io.BytesIO(base64.b64decode(response.json()['base64'])))
        self.assertEqual([p.text for p in doc.paragraphs], ['Jonathan Deng','Experience','Built APIs'])
        response = await self.client.post('/resume-document', json={'resume':' '})
        self.assertEqual(response.status_code, 400)

    async def test_tailor_missing_key_and_validation(self):
        with patch.dict(os.environ, {}, clear=True):
            response = await self.client.post('/tailor', json={'resume':'Resume','job_description':'Job'})
            self.assertEqual(response.status_code, 503)
            response = await self.client.post('/tailor', json={'resume':' ','job_description':'Job'})
            self.assertEqual(response.status_code, 400)
            response = await self.client.post('/tailor', json={'resume':'x'*40001,'job_description':'Job'})
            self.assertEqual(response.status_code, 422)

    async def test_tailor_structured_result_and_incomplete_response(self):
        parsed = TailoredResume(resume='Reviewed draft', changes=[], questions=['Confirm availability'])
        for completed in (True, False):
            with self.subTest(completed=completed):
                client = AsyncMock()
                client.__aenter__.return_value = client
                client.responses.parse.return_value = SimpleNamespace(status='completed' if completed else 'incomplete',output_parsed=parsed)
                with patch.dict(os.environ, {'JOBMATCH_OPENAI_API_KEY':'test-only'}, clear=True), patch('backend.apply.openai.AsyncOpenAI', return_value=client):
                    response = await self.client.post('/tailor', json={'resume':'Original resume','job_description':'Backend job'})
                self.assertEqual(response.status_code, 200 if completed else 502)
                args = client.responses.parse.call_args.kwargs
                self.assertIs(args['text_format'], TailoredResume)
                self.assertIn('Original resume', args['input'])
                self.assertIn('Backend job', args['input'])
                self.assertIn('Never invent', args['instructions'])
                if completed:self.assertEqual(response.json()['questions'], ['Confirm availability'])


    async def test_openai_quota_and_rate_limit_errors_are_actionable(self):
        for code in ('insufficient_quota', 'rate_limit_exceeded'):
            with self.subTest(code=code):
                client = AsyncMock()
                client.__aenter__.return_value = client
                client.responses.parse.side_effect = openai.RateLimitError(
                    'Limited', response=httpx.Response(429, request=httpx.Request('POST', 'https://api.openai.com/v1/responses')),
                    body={'code':code, 'message':'Limited', 'type':code})
                with patch.dict(os.environ, {'JOBMATCH_OPENAI_API_KEY':'test-only'}, clear=True), patch('backend.apply.openai.AsyncOpenAI', return_value=client):
                    response = await self.client.post('/tailor', json={'resume':'Original', 'job_description':'Backend role'})
                self.assertEqual(response.status_code, 429)
                self.assertIn('without AI', response.json()['detail'])
                if code == 'insufficient_quota':
                    self.assertIn('credits or quota are exhausted', response.json()['detail'])


class SavedCandidateTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        import tempfile
        self.directory = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {'JOBMATCH_DATA_DIR': self.directory.name})
        self.env.start()
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://127.0.0.1:8000')
        self.candidate = {'profile': {'firstName': 'Alex', 'email': 'alex@example.test'},
                          'resume': 'Base resume with Python APIs', 'resumeName': 'resume.txt'}

    async def asyncTearDown(self):
        await self.client.aclose()
        self.env.stop()
        self.directory.cleanup()

    async def test_persistence_is_available_to_new_sessions_and_can_be_deleted(self):
        response = await self.client.put('/candidate', json=self.candidate)
        self.assertTrue(response.json()['saved'])
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://127.0.0.1:8000') as new_session:
            saved = (await new_session.get('/candidate')).json()
        self.assertTrue(saved['ready'])
        self.assertEqual(saved['resume'], self.candidate['resume'])
        self.assertEqual(saved['profile']['email'], 'alex@example.test')
        self.assertEqual((await self.client.delete('/candidate')).status_code, 200)
        self.assertFalse((await self.client.get('/candidate')).json()['ready'])

    async def test_each_job_uses_the_base_resume_without_replacing_it(self):
        await self.client.put('/candidate', json=self.candidate)
        result = TailoredResume(resume='Tailored resume', changes=[], questions=['Check dates'])
        with patch('backend.apply.tailor', new=AsyncMock(return_value=result)) as tailor:
            for url in ['https://jobs.lever.co/example/123', 'https://job-boards.greenhouse.io/example/jobs/456']:
                response = await self.client.post('/autofill', json={'job_description':'Python engineering role with API development and teamwork. '*3, 'job_url':url})
                self.assertEqual(response.status_code, 200)
                bundle = response.json()
                self.assertTrue(bundle['generatedForJob'])
                self.assertNotIn('approved', bundle)
                self.assertEqual(bundle['profile'], self.candidate['profile'])
                doc = Document(io.BytesIO(base64.b64decode(bundle['document']['base64'])))
                self.assertEqual(doc.paragraphs[0].text, 'Tailored resume')
                self.assertEqual(tailor.call_args.args[0].resume, self.candidate['resume'])
        self.assertEqual((await self.client.get('/candidate')).json()['resume'], self.candidate['resume'])

    async def test_autofill_requires_saved_resume_and_valid_job(self):
        response = await self.client.post('/autofill', json={'job_description':'job '*30})
        self.assertEqual(response.status_code, 409)
        await self.client.put('/candidate', json=self.candidate)
        for url in ['https://evil.example/jobs/1', 'https://jobs.lever.co.evil.example/1', 'http://jobs.lever.co/1', 'https://name:password@jobs.lever.co/1']:
            response = await self.client.post('/autofill', json={'job_description':'job '*30, 'job_url':url})
            self.assertEqual(response.status_code, 400)
        response = await self.client.post('/autofill', json={'job_description':'short'})
        self.assertEqual(response.status_code, 422)

    async def test_other_websites_cannot_access_private_data_or_start_autofill(self):
        await self.client.put('/candidate', json=self.candidate)
        for method, path, payload in [('GET','/candidate',None), ('PUT','/candidate',self.candidate),
                                      ('DELETE','/candidate',None), ('POST','/autofill',{'job_description':'job '*30})]:
            response = await self.client.request(method, path, json=payload, headers={'Origin':'https://evil.example'})
            self.assertEqual(response.status_code, 403)
        response = await self.client.get('/candidate', headers={'Host':'rebinding.example'})
        self.assertEqual(response.status_code, 403)
        response = await self.client.get('/candidate', headers={'Origin':'null'})
        self.assertEqual(response.status_code, 403)
        response = await self.client.get('/candidate', headers={'Origin':'chrome-extension://'+'a'*32})
        self.assertEqual(response.status_code, 200)

    async def test_corrupt_profile_is_reported_without_overwriting_it(self):
        from backend.candidate import candidate_path
        candidate_path().write_text('broken json')
        response = await self.client.get('/candidate')
        self.assertEqual(response.status_code, 503)
        self.assertEqual(candidate_path().read_text(), 'broken json')


    async def test_oversized_generated_resume_is_not_attached_or_saved(self):
        await self.client.put('/candidate', json=self.candidate)
        result = TailoredResume(resume='x' * 40001, changes=[], questions=[])
        with patch('backend.apply.tailor', new=AsyncMock(return_value=result)):
            response = await self.client.post('/autofill', json={'job_description':'job '*30})
        self.assertEqual(response.status_code, 502)
        self.assertEqual((await self.client.get('/candidate')).json()['resume'], self.candidate['resume'])


    async def test_saved_information_autofill_never_calls_openai(self):
        await self.client.put('/candidate', json=self.candidate)
        with patch('backend.apply.openai.AsyncOpenAI') as client:
            response = await self.client.post('/autofill/saved', json={'job_url':'https://jobs.lever.co/example/1'})
        self.assertEqual(response.status_code, 200)
        client.assert_not_called()
        result = response.json()
        self.assertTrue(result['savedResume'])
        self.assertEqual(result['resume'], self.candidate['resume'])
        document = Document(io.BytesIO(base64.b64decode(result['document']['base64'])))
        self.assertEqual(document.paragraphs[0].text, self.candidate['resume'])

    async def test_contact_only_autofill_needs_neither_a_resume_nor_an_api_key(self):
        await self.client.put('/candidate', json={'profile':{'firstName':'Alex'},'resume':''})
        with patch('backend.apply.openai.AsyncOpenAI') as client:
            response = await self.client.post('/autofill/saved', json={'job_url':'https://job-boards.greenhouse.io/example/jobs/1'})
        client.assert_not_called()
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()['document'])
        self.assertEqual(response.json()['profile']['firstName'], 'Alex')

    async def test_saved_autofill_rejects_empty_profile_and_untrusted_sites(self):
        self.assertEqual((await self.client.post('/autofill/saved', json={})).status_code, 409)
        await self.client.put('/candidate', json=self.candidate)
        self.assertEqual((await self.client.post('/autofill/saved', json={'job_url':'https://evil.example'})).status_code, 400)
        self.assertEqual((await self.client.post('/autofill/saved', json={}, headers={'Origin':'https://evil.example'})).status_code, 403)
